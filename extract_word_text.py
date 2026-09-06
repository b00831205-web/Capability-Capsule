from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree


def iter_blocks(parent):
    element = parent.element.body
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def extract(path: Path):
    print(f"===== {path} =====")
    if path.suffix.lower() == ".docm":
        extract_openxml(path)
        return
    doc = Document(str(path))
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                print(text)
        else:
            for row in block.rows:
                cells = [" ".join(p.text.strip() for p in c.paragraphs if p.text.strip()) for c in row.cells]
                print(" | ".join(cells))
    for section_no, section in enumerate(doc.sections, 1):
        for label, container in (("HEADER", section.header), ("FOOTER", section.footer)):
            texts = [p.text.strip() for p in container.paragraphs if p.text.strip()]
            if texts:
                print(f"[{label} {section_no}] " + " | ".join(texts))


def extract_openxml(path: Path):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        names = ["word/document.xml"]
        names += sorted(n for n in archive.namelist() if n.startswith("word/header") and n.endswith(".xml"))
        names += sorted(n for n in archive.namelist() if n.startswith("word/footer") and n.endswith(".xml"))
        for name in names:
            root = etree.fromstring(archive.read(name))
            if name != "word/document.xml":
                print(f"[{name}]")
            for para in root.xpath(".//w:p", namespaces=ns):
                parts = para.xpath(".//w:t/text() | .//w:tab", namespaces=ns)
                text_parts = []
                for part in parts:
                    text_parts.append("\t" if not isinstance(part, str) else part)
                text = "".join(text_parts).strip()
                if text:
                    print(text)


for arg in sys.argv[1:]:
    extract(Path(arg))
